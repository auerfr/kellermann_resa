import csv
import io
import json
from datetime import date
from django.shortcuts import render, get_object_or_404
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.utils import timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import timedelta
from temple_project.apps.reservations.models import (
    Reservation, Temple, DemandeAccesPortail, ValidationSaison,
)


@login_required
def export_csv(request):
    """Export CSV des réservations selon filtres"""
    qs = _get_queryset_from_request(request)
    response = HttpResponse(content_type="text/csv; charset=utf-8-sig")
    response["Content-Disposition"] = 'attachment; filename="reservations.csv"'

    writer = csv.writer(response)
    writer.writerow([
        "Date", "Heure début", "Heure fin", "Loge", "Obédience",
        "Temple", "Type", "Sous-type", "Statut", "Agapes", "Nb repas",
        "Demandeur", "Email"
    ])
    for r in qs:
        writer.writerow([
            r.date, r.heure_debut, r.heure_fin,
            r.loge.nom, r.loge.obedience.nom, str(r.temple),
            r.get_type_reservation_display(), r.get_sous_type_display(),
            r.get_statut_display(), "Oui" if r.besoin_agapes else "Non",
            r.nombre_repas, r.nom_demandeur, r.email_demandeur,
        ])
    return response


@login_required
def export_excel(request):
    """Export Excel avec mise en forme"""
    qs = _get_queryset_from_request(request)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Réservations"

    headers = [
        "Date", "Heure début", "Heure fin", "Loge", "Obédience",
        "Temple", "Type", "Sous-type", "Statut", "Agapes", "Nb repas",
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2C3E50")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    STATUT_COLORS = {
        "validee": "C8E6C9",
        "attente": "FFF9C4",
        "refusee": "FFCDD2",
    }

    for row_idx, r in enumerate(qs, 2):
        data = [
            r.date, str(r.heure_debut), str(r.heure_fin),
            r.loge.nom, r.loge.obedience.nom, str(r.temple),
            r.get_type_reservation_display(), r.get_sous_type_display(),
            r.get_statut_display(), "Oui" if r.besoin_agapes else "Non",
            r.nombre_repas,
        ]
        fill_color = STATUT_COLORS.get(r.statut, "FFFFFF")
        fill = PatternFill("solid", fgColor=fill_color)
        for col_idx, value in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill

    # Ajustement largeur colonnes
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="reservations.xlsx"'
    return response


@login_required
def reporting(request):
    """Page de reporting et statistiques"""
    from django.db.models import Count, Sum
    from datetime import datetime

    # Saison : sept. annee → juin annee+1
    current_year = date.today().year
    # Déduire la saison courante : si on est avant septembre, la saison a démarré l'année précédente
    default_saison = current_year if date.today().month >= 9 else current_year - 1
    annee_saison = int(request.GET.get("annee", default_saison))

    saison_debut = date(annee_saison, 9, 1)
    saison_fin   = date(annee_saison + 1, 6, 30)

    reservations = Reservation.objects.filter(date__gte=saison_debut, date__lte=saison_fin)

    total      = reservations.count()
    validees   = reservations.filter(statut="validee").count()
    attente    = reservations.filter(statut="attente").count()
    refusees   = reservations.filter(statut="refusee").count()
    total_repas = reservations.filter(besoin_agapes=True, statut="validee").aggregate(
        s=Sum("nombre_repas")
    )["s"] or 0

    stats = {
        "total": total,
        "validees": validees,
        "attente": attente,
        "refusees": refusees,
        "total_repas": total_repas,
        "taux_validation": round(validees / total * 100, 1) if total > 0 else 0,
    }

    # Réservations par obédience (toute la saison, tous statuts)
    reservations_par_obedience = reservations.values(
        'loge__obedience__nom'
    ).annotate(
        nb_reservations=Count('id')
    ).order_by('-nb_reservations')[:10]

    # Graphique mensuel : mois de la saison (sept → juin), tous statuts
    reservations_par_mois = []
    mois_saison = [9, 10, 11, 12, 1, 2, 3, 4, 5, 6]
    for m in mois_saison:
        annee_mois = annee_saison if m >= 9 else annee_saison + 1
        count = reservations.filter(date__year=annee_mois, date__month=m).count()
        reservations_par_mois.append({
            'mois': f'{annee_mois}-{m:02d}',
            'count': count,
        })

    # Réservations par temple (tous statuts)
    reservations_par_temple = reservations.values(
        'temple__nom'
    ).annotate(
        nb_reservations=Count('id')
    ).order_by('-nb_reservations')

    reservations_par_mois_json = json.dumps(reservations_par_mois)
    reservations_par_temple_json = json.dumps([
        {'nom': t['temple__nom'] or 'Non renseigné', 'nb_reservations': t['nb_reservations']}
        for t in reservations_par_temple
    ])

    context = {
        "stats": stats,
        "annee": annee_saison,
        "annee_courante": default_saison,
        "saison_label": f"{annee_saison}–{annee_saison + 1}",
        "reservations_par_obedience": reservations_par_obedience,
        "reservations_par_mois": reservations_par_mois_json,
        "reservations_par_temple": reservations_par_temple_json,
        "temples": Temple.objects.all().order_by('nom'),
    }

    return render(request, "exports/reporting.html", context)


def _get_queryset_from_request(request):
    qs = Reservation.objects.select_related("loge", "loge__obedience", "temple")
    if request.GET.get("mois"):
        qs = qs.filter(date__month=request.GET["mois"])
    if request.GET.get("annee"):
        qs = qs.filter(date__year=request.GET["annee"])
    if request.GET.get("temple"):
        qs = qs.filter(temple_id=request.GET["temple"])
    if request.GET.get("loge"):
        qs = qs.filter(loge_id=request.GET["loge"])
    return qs.order_by("date", "heure_debut")


@login_required
def planning_pdf(request):
    """Export PDF du planning des tenues, par temple — modes mois / saison / perso."""
    from io import BytesIO
    import calendar as _cal
    from django.db.models import Sum
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle,
        Paragraph, Spacer, HRFlowable, PageBreak,
    )
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    today    = date.today()
    temple_p = request.GET.get('temple') or None
    mode     = request.GET.get('mode', 'mois')

    MOIS_NOMS = {1: 'Janvier', 2: 'Février', 3: 'Mars', 4: 'Avril', 5: 'Mai',
                 6: 'Juin', 7: 'Juillet', 8: 'Août', 9: 'Septembre',
                 10: 'Octobre', 11: 'Novembre', 12: 'Décembre'}
    JOURS_FR  = ['Lun', 'Mar', 'Mer', 'Jeu', 'Ven', 'Sam', 'Dim']

    # ── Période selon le mode ──────────────────────────────────────────────────
    if mode == 'saison':
        default_annee = today.year if today.month >= 9 else today.year - 1
        annee_p       = int(request.GET.get('annee', default_annee))
        debut         = date(annee_p, 9, 1)
        fin           = date(annee_p + 1, 6, 30)
        titre_periode = f"Saison {annee_p}\u2013{annee_p + 1}"
        nom_fichier   = f"planning_saison_{annee_p}-{annee_p + 1}"
    elif mode == 'perso':
        try:
            debut = date.fromisoformat(request.GET.get('date_debut', str(today)))
            fin   = date.fromisoformat(request.GET.get('date_fin',   str(today)))
        except ValueError:
            debut = fin = today
        annee_p       = debut.year
        titre_periode = (f"Du {debut.strftime('%d/%m/%Y')} "
                         f"au {fin.strftime('%d/%m/%Y')}")
        nom_fichier   = f"planning_{debut}_{fin}"
    else:  # mois (défaut)
        mois_p  = int(request.GET.get('mois',  today.month))
        annee_p = int(request.GET.get('annee', today.year))
        debut   = date(annee_p, mois_p, 1)
        fin     = date(annee_p, mois_p, _cal.monthrange(annee_p, mois_p)[1])
        titre_periode = f"{MOIS_NOMS[mois_p]} {annee_p}"
        nom_fichier   = f"planning_{annee_p}_{mois_p:02d}"

    if temple_p:
        temples = list(Temple.objects.filter(pk=temple_p))
    else:
        temples = list(Temple.objects.all().order_by('nom'))

    C_NAVY  = colors.HexColor('#0F2137')
    C_GOLD  = colors.HexColor('#C8A84B')
    C_LIGHT = colors.HexColor('#F8FAFC')

    # Couleurs de badge pour la colonne Type
    TYPE_BG = {
        'reguliere':      colors.HexColor('#DBEAFE'),
        'exceptionnelle': colors.HexColor('#FEF3C7'),
        'congres':        colors.HexColor('#FCE7F3'),
    }

    sty_h1   = ParagraphStyle('pg_h1',  fontName='Helvetica-Bold', fontSize=14,
                               textColor=C_NAVY, spaceAfter=2)
    sty_h2   = ParagraphStyle('pg_h2',  fontName='Helvetica-Bold', fontSize=11,
                               textColor=C_GOLD, spaceAfter=6)
    sty_sub  = ParagraphStyle('pg_sub', fontName='Helvetica',      fontSize=9,
                               textColor=colors.grey, spaceAfter=10)
    sty_vide = ParagraphStyle('pg_vide', fontName='Helvetica-Oblique', fontSize=9,
                               textColor=colors.grey, alignment=TA_CENTER)

    buf = BytesIO()

    def _footer(canvas, doc):
        if doc.page == 1:
            return
        canvas.saveState()
        canvas.setFont('Helvetica', 7)
        canvas.setFillColorRGB(0.55, 0.55, 0.55)
        page_num = doc.page - 1
        canvas.drawCentredString(
            A4[0] / 2, 1.2 * cm,
            f"Kellermann \u00b7 Planning des tenues \u00b7 "
            f"{titre_periode} \u00b7 Page {page_num}",
        )
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm,  bottomMargin=2.2 * cm,
    )

    col_widths = [2.2 * cm, 1.4 * cm, 4.8 * cm, 2.8 * cm, 2.6 * cm, 2.0 * cm]
    headers    = ['Date', 'Jour', 'Loge', 'Horaires', 'Type', 'Agapes']

    def _table_temple(temple):
        tenues = (
            Reservation.objects
            .select_related('loge', 'loge__obedience', 'temple')
            .filter(temple=temple, statut='validee',
                    date__gte=debut, date__lte=fin)
            .order_by('date', 'heure_debut')
        )
        rows        = [headers]
        type_colors = []
        for i, t in enumerate(tenues, 1):
            loge_nom = (t.loge.nom if t.loge
                        else (t.nom_organisation
                              if hasattr(t, 'nom_organisation') else '\u2014'))
            agapes = (f"\u2713 {t.nombre_repas} cvts"
                      if t.besoin_agapes else '\u2014')
            rows.append([
                t.date.strftime('%d/%m/%Y'),
                JOURS_FR[t.date.weekday()],
                loge_nom,
                f"{t.heure_debut:%H:%M} \u2013 {t.heure_fin:%H:%M}",
                t.get_type_reservation_display(),
                agapes,
            ])
            if t.type_reservation in TYPE_BG:
                type_colors.append((i, TYPE_BG[t.type_reservation]))
        return rows, tenues.count(), type_colors

    story = []

    # ── Page de garde ──────────────────────────────────────────────────────────
    story.append(Spacer(1, 3 * cm))
    story.append(Paragraph('\u2692', ParagraphStyle(
        'gd_ico', fontName='Helvetica-Bold', fontSize=38,
        textColor=C_NAVY, alignment=TA_CENTER, spaceAfter=4)))
    story.append(Paragraph('Temples Kellermann', ParagraphStyle(
        'gd_brand', fontName='Helvetica-Bold', fontSize=13,
        textColor=C_GOLD, alignment=TA_CENTER, spaceAfter=16)))
    story.append(HRFlowable(width='50%', thickness=0.8, color=C_GOLD, hAlign='CENTER'))
    story.append(Spacer(1, 1 * cm))
    story.append(Paragraph('Planning des tenues', ParagraphStyle(
        'gd_titre', fontName='Helvetica-Bold', fontSize=22,
        textColor=C_NAVY, alignment=TA_CENTER, spaceAfter=8)))
    story.append(Paragraph(titre_periode, ParagraphStyle(
        'gd_periode', fontName='Helvetica', fontSize=14,
        textColor=colors.HexColor('#555555'), alignment=TA_CENTER, spaceAfter=6)))
    story.append(Paragraph(
        f'G\u00e9n\u00e9r\u00e9 le {today.strftime("%d/%m/%Y")}',
        ParagraphStyle('gd_gen', fontName='Helvetica', fontSize=9,
                       textColor=colors.grey, alignment=TA_CENTER)))
    story.append(PageBreak())

    # ── En-tête de contenu ─────────────────────────────────────────────────────
    story.append(Paragraph("Kellermann \u2014 Planning des tenues", sty_h1))
    story.append(Paragraph(titre_periode, sty_h2))
    story.append(Spacer(1, 0.2 * cm))

    # ── Tables par temple ──────────────────────────────────────────────────────
    for i, temple in enumerate(temples):
        rows, nb, type_colors = _table_temple(temple)

        story.append(Paragraph(str(temple), ParagraphStyle(
            f'tnom_{i}', fontName='Helvetica-Bold', fontSize=10,
            textColor=C_NAVY, spaceAfter=3,
        )))
        story.append(Paragraph(
            f"{nb} tenue{'s' if nb != 1 else ''}", sty_sub))

        if nb == 0:
            story.append(Paragraph("Aucune tenue sur cette p\u00e9riode.", sty_vide))
        else:
            n = len(rows)
            style_cmds = [
                ('BACKGROUND',    (0, 0),  (-1, 0),   C_NAVY),
                ('TEXTCOLOR',     (0, 0),  (-1, 0),   C_GOLD),
                ('FONTNAME',      (0, 0),  (-1, 0),   'Helvetica-Bold'),
                ('FONTSIZE',      (0, 0),  (-1, 0),   8),
                ('ALIGN',         (0, 0),  (-1, 0),   'CENTER'),
                ('TOPPADDING',    (0, 0),  (-1, 0),   5),
                ('BOTTOMPADDING', (0, 0),  (-1, 0),   5),
                ('FONTNAME',      (0, 1),  (-1, n-1), 'Helvetica'),
                ('FONTSIZE',      (0, 1),  (-1, n-1), 8),
                ('ROWBACKGROUNDS',(0, 1),  (-1, n-1), [colors.white, C_LIGHT]),
                ('TOPPADDING',    (0, 1),  (-1, n-1), 4),
                ('BOTTOMPADDING', (0, 1),  (-1, n-1), 4),
                ('ALIGN',         (1, 1),  (1, n-1),  'CENTER'),
                ('ALIGN',         (3, 1),  (3, n-1),  'CENTER'),
                ('ALIGN',         (5, 1),  (5, n-1),  'CENTER'),
                ('GRID',          (0, 0),  (-1, -1),  0.4, colors.HexColor('#CBD5E1')),
                ('BOX',           (0, 0),  (-1, -1),  1,   C_NAVY),
            ]
            # Badges colorés par type (colonne 4)
            for ri, bg in type_colors:
                style_cmds.append(('BACKGROUND', (4, ri), (4, ri), bg))

            tbl = Table(rows, colWidths=col_widths, repeatRows=1)
            tbl.setStyle(TableStyle(style_cmds))
            story.append(tbl)

        if i < len(temples) - 1:
            story.append(Spacer(1, 0.5 * cm))
            story.append(HRFlowable(width='100%', thickness=0.5,
                                    color=colors.HexColor('#E2E8F0')))
            story.append(Spacer(1, 0.3 * cm))

    # ── Résumé global ──────────────────────────────────────────────────────────
    story.append(Spacer(1, 0.6 * cm))
    story.append(HRFlowable(width='100%', thickness=0.5,
                             color=colors.HexColor('#E2E8F0')))
    story.append(Spacer(1, 0.3 * cm))

    filter_kw = dict(statut='validee', date__gte=debut, date__lte=fin)
    if temple_p:
        filter_kw['temple__pk'] = temple_p
    summary_qs   = Reservation.objects.filter(**filter_kw)
    total_tenues = summary_qs.count()
    total_cvts   = (summary_qs.filter(besoin_agapes=True)
                               .aggregate(s=Sum('nombre_repas'))['s'] or 0)

    for t in temples:
        nb_t = summary_qs.filter(temple=t).count()
        if nb_t:
            story.append(Paragraph(
                f"{t}\u00a0: {nb_t} tenue(s)",
                ParagraphStyle(f'sl_{t.pk}', fontName='Helvetica', fontSize=8,
                               textColor=colors.HexColor('#555555'), spaceAfter=2),
            ))
    story.append(Spacer(1, 0.2 * cm))
    story.append(Paragraph(
        f"Total\u00a0: {total_tenues} tenue(s) \u00b7 {total_cvts} couverts agapes",
        ParagraphStyle('pg_total', fontName='Helvetica-Bold', fontSize=8,
                       textColor=C_NAVY, spaceAfter=0),
    ))

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    buf.seek(0)
    if temple_p and temples:
        nom_fichier += f"_{temples[0].nom.replace(' ', '_')}"
    nom_fichier += ".pdf"
    response = HttpResponse(buf, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{nom_fichier}"'
    return response


def planning_loge_pdf(request, token):
    """Export PDF du planning de saison d'une loge, accessible via token portail."""
    from io import BytesIO
    from django.db.models import Sum
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle,
        Paragraph, Spacer, HRFlowable, PageBreak,
    )
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    demande = get_object_or_404(DemandeAccesPortail, token=token, statut='validee')
    loge    = demande.loge
    if not loge:
        return HttpResponse("Aucune loge associée à ce token.", status=400)

    today = date.today()

    # Priorité : 1) paramètre GET ?annee=  2) ValidationSaison active  3) today
    annee_default = today.year if today.month >= 9 else today.year - 1
    if request.GET.get('annee'):
        try:
            annee = int(request.GET['annee'])
        except ValueError:
            annee = annee_default
    else:
        val = ValidationSaison.objects.filter(
            loge=loge,
            statut__in=['ouverte', 'soumise'],
        ).order_by('-annee').first()
        annee = val.annee if val else annee_default

    debut_saison = date(annee, 9, 1)
    fin_saison   = date(annee + 1, 6, 30)

    tenues = (
        Reservation.objects
        .select_related('temple')
        .filter(loge=loge, statut='validee',
                date__gte=debut_saison, date__lte=fin_saison)
        .order_by('date', 'heure_debut')
    )

    JOURS_FR = ['Lun', 'Mar', 'Mer', 'Jeu', 'Ven', 'Sam', 'Dim']
    C_HEADER = colors.HexColor('#222222')
    C_LIGHT  = colors.HexColor('#F2F2F2')
    C_BORDER = colors.HexColor('#BBBBBB')

    buf = BytesIO()

    def _footer(canvas, doc):
        if doc.page == 1:
            return
        canvas.saveState()
        canvas.setFont('Helvetica', 7)
        canvas.setFillColorRGB(0.55, 0.55, 0.55)
        canvas.drawCentredString(
            A4[0] / 2, 1.2 * cm,
            f"Temples Kellermann \u00b7 {loge.nom} \u00b7 "
            f"Saison {annee}\u2013{annee + 1} \u00b7 Page {doc.page - 1}",
        )
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=1.8 * cm, rightMargin=1.8 * cm,
        topMargin=1.5 * cm,  bottomMargin=2.2 * cm,
    )

    def _sty(name, size, bold=False, color='#0F2137', align=TA_CENTER, after=6):
        return ParagraphStyle(name,
                              fontName='Helvetica-Bold' if bold else 'Helvetica',
                              fontSize=size,
                              textColor=colors.HexColor(color),
                              alignment=align,
                              spaceAfter=after)

    story = []

    # ── Page de garde ──────────────────────────────────────────────────────────
    story.append(Spacer(1, 3 * cm))
    story.append(Paragraph('\u2692', _sty('lp_ico', 38, bold=True)))
    story.append(Paragraph('Temples Kellermann',
                            _sty('lp_brand', 13, bold=True, color='#C8A84B', after=16)))
    story.append(HRFlowable(width='50%', thickness=0.8,
                             color=colors.HexColor('#C8A84B'), hAlign='CENTER'))
    story.append(Spacer(1, 1 * cm))
    story.append(Paragraph('Planning des tenues', _sty('lp_titre', 20, bold=True)))
    story.append(Spacer(1, 0.5 * cm))
    story.append(Paragraph(loge.nom, _sty('lp_loge', 14, bold=True, color='#333333')))
    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph(f'Saison {annee}\u2013{annee + 1}',
                            _sty('lp_saison', 11, color='#555555')))
    story.append(Paragraph(f'01/09/{annee} \u2192 30/06/{annee + 1}',
                            _sty('lp_periode', 9, color='#888888')))
    story.append(Spacer(1, 2.5 * cm))
    story.append(Paragraph(f'G\u00e9n\u00e9r\u00e9 le {today.strftime("%d/%m/%Y")}',
                            _sty('lp_gen', 8, color='#999999')))
    story.append(PageBreak())

    # ── Tableau des tenues ─────────────────────────────────────────────────────
    nb = tenues.count()
    story.append(Paragraph(
        f'Planning saison {annee}\u2013{annee + 1}',
        _sty('lt_h2', 11, bold=True, align=TA_LEFT, after=3),
    ))
    story.append(Paragraph(
        f'{nb} tenue{"s" if nb != 1 else ""} valid\u00e9e{"s" if nb != 1 else ""}',
        _sty('lt_sub', 9, color='#888888', align=TA_LEFT, after=8),
    ))

    if nb == 0:
        story.append(Paragraph(
            'Aucune tenue valid\u00e9e pour cette saison.',
            _sty('lt_vide', 9, color='#888888'),
        ))
    else:
        col_widths = [2.2 * cm, 1.3 * cm, 3.8 * cm, 2.8 * cm, 3.8 * cm, 2.5 * cm]
        rows = [['Date', 'Jour', 'Temple', 'Horaires', 'Type', 'Agapes']]
        for t in tenues:
            agapes = f"\u2713 {t.nombre_repas} cvts" if t.besoin_agapes else '\u2014'
            rows.append([
                t.date.strftime('%d/%m/%Y'),
                JOURS_FR[t.date.weekday()],
                str(t.temple) if t.temple else '\u2014',
                f"{t.heure_debut:%H:%M} \u2013 {t.heure_fin:%H:%M}",
                t.get_type_reservation_display(),
                agapes,
            ])
        nr = len(rows)
        tbl = Table(rows, colWidths=col_widths, repeatRows=1)
        tbl.setStyle(TableStyle([
            ('BACKGROUND',    (0, 0),   (-1, 0),    C_HEADER),
            ('TEXTCOLOR',     (0, 0),   (-1, 0),    colors.white),
            ('FONTNAME',      (0, 0),   (-1, 0),    'Helvetica-Bold'),
            ('FONTSIZE',      (0, 0),   (-1, -1),   8),
            ('ALIGN',         (0, 0),   (-1, 0),    'CENTER'),
            ('TOPPADDING',    (0, 0),   (-1, -1),   4),
            ('BOTTOMPADDING', (0, 0),   (-1, -1),   4),
            ('FONTNAME',      (0, 1),   (-1, nr-1), 'Helvetica'),
            ('ROWBACKGROUNDS',(0, 1),   (-1, nr-1), [colors.white, C_LIGHT]),
            ('ALIGN',         (1, 1),   (1, nr-1),  'CENTER'),
            ('ALIGN',         (3, 1),   (3, nr-1),  'CENTER'),
            ('ALIGN',         (5, 1),   (5, nr-1),  'CENTER'),
            ('GRID',          (0, 0),   (-1, -1),   0.3, C_BORDER),
            ('BOX',           (0, 0),   (-1, -1),   0.6, colors.HexColor('#666666')),
        ]))
        story.append(tbl)

        nb_agapes  = tenues.filter(besoin_agapes=True).count()
        total_cvts = (tenues.filter(besoin_agapes=True)
                             .aggregate(s=Sum('nombre_repas'))['s'] or 0)
        story.append(Spacer(1, 0.4 * cm))
        story.append(Paragraph(
            f'Total\u00a0: {nb} tenue(s)  \u00b7  '
            f'{nb_agapes} avec agapes ({total_cvts} couverts)',
            _sty('lt_resume', 8, color='#777777', align=TA_LEFT, after=0),
        ))

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    buf.seek(0)
    nom = f"planning_{loge.nom.replace(' ', '_')}_{annee}-{annee + 1}.pdf"
    response = HttpResponse(buf, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{nom}"'
    return response
